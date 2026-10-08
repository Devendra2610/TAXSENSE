import os
import sys
import pypdf
import docx

sys.stdout.reconfigure(encoding='utf-8')

data_dir = r"C:\Users\hp\Desktop\CA\DATA"

def dump_pdf(path, out_path):
    try:
        reader = pypdf.PdfReader(path)
        with open(out_path, "w", encoding="utf-8") as f:
            for i, page in enumerate(reader.pages):
                f.write(f"\n--- PAGE {i} ---\n")
                f.write(page.extract_text() or "")
        print(f"Dumped PDF to {out_path}")
    except Exception as e:
        print(f"Error dumping PDF: {e}")

def dump_docx(path, out_path):
    try:
        doc = docx.Document(path)
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("=== PARAGRAPHS ===\n")
            for p in doc.paragraphs:
                if p.text.strip():
                    f.write(p.text + "\n")
            f.write("\n=== TABLES ===\n")
            for t_idx, table in enumerate(doc.tables):
                f.write(f"\n--- TABLE {t_idx} ---\n")
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    f.write(" | ".join(cells) + "\n")
        print(f"Dumped DOCX to {out_path}")
    except Exception as e:
        print(f"Error dumping DOCX: {e}")

dump_pdf(os.path.join(data_dir, "Comp.pdf"), r"C:\Users\hp\Desktop\CA\comp_text.txt")
dump_pdf(os.path.join(data_dir, "sub report.pdf"), r"C:\Users\hp\Desktop\CA\sub_report_text.txt")
dump_docx(os.path.join(data_dir, "TAX AUDIT REPORT FY 2025-26.docx"), r"C:\Users\hp\Desktop\CA\tax_audit_report_text.txt")
